"""Load cedant names from the client's Excel file into ``cedant`` (issue 129).

Alembic revision 0005 and ``infra/scripts/load_cedants.py`` both call it, so it
imports nothing from ``app``.
"""

from __future__ import annotations

import uuid
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import text

SEED_FILE = Path(__file__).resolve().parent / "bootstrap" / "seed" / "cedants.xlsx"
HEADER = "Cedant"
MAX_NAME_LENGTH = 255  # cedant.name is NVARCHAR(255)


def read_cedant_names(path: Path) -> list[str]:
    """The names in the first sheet's ``Cedant`` column, trimmed, blanks
    skipped, and repeats (case ignored) dropped in favour of the first spelling.

    Raises ``ValueError`` when the first cell is not ``Cedant`` or a name is
    longer than 255 characters, naming the Excel row.
    """
    workbook = load_workbook(path, read_only=True)
    try:
        rows = workbook.worksheets[0].iter_rows(max_col=1, values_only=True)
        header = next(rows, ())
        if not header or str(header[0] or "").strip().lower() != HEADER.lower():
            raise ValueError(f"{path.name}: cell A1 must read {HEADER!r}")
        names: dict[str, str] = {}
        for row_number, row in enumerate(rows, start=2):
            name = str(row[0]).strip() if row and row[0] is not None else ""
            if not name:
                continue
            if len(name) > MAX_NAME_LENGTH:
                raise ValueError(f"{path.name} row {row_number}: the cedant name is "
                                 f"longer than {MAX_NAME_LENGTH} characters")
            names.setdefault(name.lower(), name)
        return list(names.values())
    finally:
        workbook.close()


def load_cedant_names(conn, names: list[str]) -> int:
    """Insert each name not already in ``cedant`` (case ignored) as an active
    cedant on the caller's transaction; return how many were inserted. Never
    renames or reactivates an existing cedant."""
    existing = {name.lower() for name in conn.execute(
        text("SELECT name FROM cedant")).scalars()}
    now = datetime.now(UTC).replace(tzinfo=None)
    rows = [{"id": str(uuid.uuid4()), "name": name, "now": now}
            for name in names if name.lower() not in existing]
    if rows:
        conn.execute(text(
            "INSERT INTO cedant (id, name, is_active, inserted_at, updated_at) "
            "VALUES (:id, :name, 1, :now, :now)"), rows)
    return len(rows)
